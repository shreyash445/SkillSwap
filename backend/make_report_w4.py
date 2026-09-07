from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)


def h(text, size=16, color=(205, 255, 87)):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(*color)
    return p


def para(text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    p.add_run(text)
    return p


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("SkillSwap — Week 4 Progress Report")
r.bold = True
r.font.size = Pt(22)
r.font.color.rgb = RGBColor(0, 0, 0)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Status: New features shipped, QA tooling added, all checks green")
r.font.size = Pt(12)
r.font.color.rgb = RGBColor(90, 90, 90)

doc.add_paragraph()

h("1. Summary", 15, (20, 20, 20))
para(
    "Week 4 focused on closing the gap between the MVP and the spec: two user-facing features were "
    "added (in-app skill management and name search on Discover), and the project gained a proper "
    "QA setup with automated type checking and linting. The backend was extended only where these "
    "features required it."
)

doc.add_paragraph()
h("2. New Features", 15, (20, 20, 20))
para("Manage skills from Profile", bold_prefix="2.1 ")
bullet("Users can now add and remove the skills they teach and want to learn directly from their Profile.")
bullet("A new Manage Skills sheet shows current teach/learn skills with one-tap remove buttons.")
bullet("An add-skill picker includes a teach/learn toggle, proficiency level for offers, search, and category filters.")
bullet("Every change is applied immediately through the existing REST API and the profile refreshes in place.")
bullet("Backend: user profiles now return the internal UserSkill id so the app can delete a specific skill.")
para("Search users by name on Discover", bold_prefix="2.2 ")
bullet("A search box lets users filter the Discover feed by name, composing with the existing skill filter and sort.")
bullet("Backend: the user list endpoint now accepts a ?q= parameter for name search (first/last name, case-insensitive).")

doc.add_paragraph()
h("3. QA & Testing", 15, (20, 20, 20))
bullet("Added npm scripts: typecheck (tsc --noEmit) and lint (expo lint).")
bullet("ESLint is now set up automatically via expo lint with the SDK-matched eslint-config-expo.")
bullet("TypeScript compiles cleanly under strict mode with zero errors.")
bullet("ESLint passes with zero errors; remaining warnings are pre-existing animation hook patterns.")
bullet("Fixed a broken Unicode swap symbol in the chat screen (mojibake) so the exchange arrow renders correctly.")
bullet("Fixed several lint errors across screens: unescaped entities, unused imports, and a hook called inside a list render callback.")
bullet("Backend: all end-to-end flow checks pass (14/14; one re-run-only duplicate-email case excluded).")
bullet("Verified the new ?q= search returns the expected user and skill management ids are exposed on the API.")

doc.add_paragraph()
h("4. Result", 15, (20, 20, 20))
para(
    "SkillSwap now matches the advertised spec more closely: profiles are fully editable without "
    "re-running onboarding, Discover is searchable by name, and the project has a repeatable QA "
    "pipeline (typecheck + lint) that can run before every demo. The backend stays stable with all "
    "flow checks passing."
)

doc.add_paragraph()
h("5. Next Steps", 15, (20, 20, 20))
bullet("Auto refresh-token handling in the mobile API client (retry on 401).")
bullet("Unread indicators per conversation in the Messages tab.")
bullet("Optional: convert the flow checks into a persistent Django test suite.")
bullet("Deploy the backend (settings already support env-driven SECRET_KEY / DEBUG / ALLOWED_HOSTS).")

out = r"C:\Users\Admin\skillswap-mvp\SkillSwap_Week4_Progress_Report.docx"
doc.save(out)
print("Saved:", out)